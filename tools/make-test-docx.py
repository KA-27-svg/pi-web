"""生成一个用于验证 docx 输出能力的测试文档。"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# 标题
h = doc.add_heading("Docx 生成测试文档", level=0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("本文件由 python-docx 生成，用于验证标题、正文、列表、表格、样式等基础能力。")

doc.add_heading("1. 纯文本与格式", level=1)
p = doc.add_paragraph("普通段落。")
p.add_run("加粗").bold = True
p.add_run("、")
p.add_run("斜体").italic = True
p.add_run("、")
r = p.add_run("彩色字")
r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
p.add_run("，以及 ")
r2 = p.add_run("等宽代码体")
r2.font.name = "Consolas"
r2.font.size = Pt(10)

doc.add_heading("2. 项目符号与编号", level=1)
for item in ["第一项", "第二项", "第三项"]:
    doc.add_paragraph(item, style="List Bullet")
for item in ["步骤一：准备数据", "步骤二：生成文档", "步骤三：校验输出"]:
    doc.add_paragraph(item, style="List Number")

doc.add_heading("3. 表格", level=1)
table = doc.add_table(rows=1, cols=3)
table.style = "Light Grid Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
for i, name in enumerate(["检查项", "预期结果", "状态"]):
    hdr[i].text = name
rows = [
    ("文档可打开", "Word/WPS 正常渲染", "通过"),
    ("中文字体", "无乱码、无缺字", "通过"),
    ("样式生效", "标题/列表/表格样式正确", "通过"),
]
for c1, c2, c3 in rows:
    cells = table.add_row().cells
    cells[0].text, cells[1].text, cells[2].text = c1, c2, c3

doc.add_heading("4. 图片占位", level=1)
doc.add_paragraph("（此处可插入图片）").alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("5. 页脚信息", level=1)
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(1)
foot = sec.footer.paragraphs[0]
foot.text = "pi web · docx 测试"
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER

out = "docx-test.docx"
doc.save(out)
print(f"saved: {out}")
